"""Document loaders for the local knowledge base."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from core.knowledge.models import SourceDocument

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".csv",
    ".tsv",
    ".pdf",
    ".docx",
}


def load_document(path: str | Path) -> SourceDocument:
    file_path = Path(path).resolve()
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported knowledge file type: {suffix}")

    metadata = {"extension": suffix}

    if suffix == ".txt":
        content = _read_text(file_path)
    elif suffix in {".md", ".markdown"}:
        content, markdown_metadata = _read_markdown(file_path)
        metadata.update(markdown_metadata)
    elif suffix == ".json":
        content = _read_json(file_path)
    elif suffix in {".csv", ".tsv"}:
        content = _read_table(file_path, delimiter="\t" if suffix == ".tsv" else ",")
    elif suffix == ".pdf":
        content = _read_pdf(file_path)
    elif suffix == ".docx":
        content = _read_docx(file_path)
    else:
        content = ""

    content = _normalize_text(content)
    if not content:
        raise ValueError(f"knowledge file has no readable text: {file_path}")
    return SourceDocument(
        source_path=str(file_path),
        title=file_path.stem,
        content=content,
        metadata=metadata,
    )


def iter_documents(paths: list[str | Path]) -> list[SourceDocument]:
    documents: list[SourceDocument] = []
    for raw_path in paths:
        path = Path(raw_path).resolve()
        if path.is_dir():
            for child in sorted(path.rglob("*")):
                if child.is_file() and child.suffix.lower() in SUPPORTED_EXTENSIONS:
                    documents.append(load_document(child))
        elif path.is_file():
            documents.append(load_document(path))
        else:
            raise FileNotFoundError(f"knowledge path not found: {path}")
    return documents


def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_markdown(path: Path) -> tuple[str, dict[str, str]]:
    text = _read_text(path)
    lines: list[str] = []
    headings: list[tuple[int, str]] = []
    heading_stack: list[tuple[int, str]] = []
    table_header: list[str] | None = None
    in_code_block = False

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```") or stripped.startswith("~~~"):
            in_code_block = not in_code_block
            table_header = None
            continue
        if in_code_block:
            continue

        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", stripped)
        if heading:
            level = len(heading.group(1))
            title = _strip_markdown_inline(heading.group(2))
            if title:
                heading_stack = [item for item in heading_stack if item[0] < level]
                heading_stack.append((level, title))
                headings.append((level, title))
                lines.append(f"{' > '.join(item[1] for item in heading_stack)}")
            table_header = None
            continue

        if _is_markdown_table_separator(stripped):
            continue

        if _looks_like_markdown_table_row(stripped):
            cells = _split_markdown_table_row(stripped)
            if not cells:
                continue
            if table_header is None:
                table_header = cells
                lines.append("表格字段：" + "，".join(table_header))
            else:
                pairs = []
                for index, value in enumerate(cells):
                    header = table_header[index] if index < len(table_header) else f"字段{index + 1}"
                    pairs.append(f"{header}：{value}")
                if pairs:
                    lines.append("；".join(pairs))
            continue

        table_header = None
        cleaned = _strip_markdown_block_prefix(stripped)
        cleaned = _strip_markdown_inline(cleaned)
        if cleaned:
            if heading_stack and not cleaned.startswith(heading_stack[-1][1]):
                section = " > ".join(item[1] for item in heading_stack)
                lines.append(f"{section}：{cleaned}")
            else:
                lines.append(cleaned)

    metadata: dict[str, str] = {"loader": "markdown"}
    if headings:
        metadata["headings"] = " > ".join(title for _, title in headings[:12])
        metadata["primary_heading"] = headings[0][1]
    return "\n".join(lines), metadata


def _strip_markdown_block_prefix(text: str) -> str:
    text = re.sub(r"^\s{0,3}>\s?", "", text)
    text = re.sub(r"^\s*-\s+\[[ xX]\]\s+", "", text)
    text = re.sub(r"^\s*[-*+]\s+", "", text)
    text = re.sub(r"^\s*\d+[.)]\s+", "", text)
    return text.strip()


def _strip_markdown_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"[*_~]{1,3}", "", text)
    text = text.replace("<br>", " ").replace("<br/>", " ").replace("<br />", " ")
    return " ".join(text.split())


def _looks_like_markdown_table_row(text: str) -> bool:
    return text.startswith("|") and text.endswith("|") and text.count("|") >= 2


def _is_markdown_table_separator(text: str) -> bool:
    if not _looks_like_markdown_table_row(text):
        return False
    cells = _split_markdown_table_row(text)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _split_markdown_table_row(text: str) -> list[str]:
    cells = [cell.strip() for cell in text.strip().strip("|").split("|")]
    return [_strip_markdown_inline(cell) for cell in cells if cell.strip()]


def _read_json(path: Path) -> str:
    data = json.loads(_read_text(path))
    return json.dumps(data, ensure_ascii=False, indent=2)


def _read_table(path: Path, delimiter: str) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter=delimiter)
        for row in reader:
            values = [item.strip() for item in row if item.strip()]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("reading PDF knowledge files requires pypdf") from exc

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[第{index}页]\n{text}")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("reading DOCX knowledge files requires python-docx") from exc

    document = docx.Document(str(path))
    lines = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _normalize_text(text: str) -> str:
    lines = [" ".join(line.strip().split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)
