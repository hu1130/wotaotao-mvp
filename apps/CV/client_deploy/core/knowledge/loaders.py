"""Document loaders for the local knowledge base."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from core.knowledge.models import SourceDocument

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".csv",
    ".tsv",
    ".pdf",
    ".docx",
}


def load_document(path: str | Path) -> SourceDocument:
    file_path = Path(path).resolve()
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported knowledge file type: {suffix}")

    if suffix in {".txt", ".md", ".markdown"}:
        content = _read_text(file_path)
    elif suffix == ".json":
        content = _read_json(file_path)
    elif suffix in {".csv", ".tsv"}:
        content = _read_table(file_path, delimiter="\t" if suffix == ".tsv" else ",")
    elif suffix == ".pdf":
        content = _read_pdf(file_path)
    elif suffix == ".docx":
        content = _read_docx(file_path)
    else:
        content = ""

    content = _normalize_text(content)
    if not content:
        raise ValueError(f"knowledge file has no readable text: {file_path}")
    return SourceDocument(
        source_path=str(file_path),
        title=file_path.stem,
        content=content,
        metadata={"extension": suffix},
    )


def iter_documents(paths: list[str | Path]) -> list[SourceDocument]:
    documents: list[SourceDocument] = []
    for raw_path in paths:
        path = Path(raw_path).resolve()
        if path.is_dir():
            for child in sorted(path.rglob("*")):
                if child.is_file() and child.suffix.lower() in SUPPORTED_EXTENSIONS:
                    documents.append(load_document(child))
        elif path.is_file():
            documents.append(load_document(path))
        else:
            raise FileNotFoundError(f"knowledge path not found: {path}")
    return documents


def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_json(path: Path) -> str:
    data = json.loads(_read_text(path))
    return json.dumps(data, ensure_ascii=False, indent=2)


def _read_table(path: Path, delimiter: str) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter=delimiter)
        for row in reader:
            values = [item.strip() for item in row if item.strip()]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("reading PDF knowledge files requires pypdf") from exc

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[第{index}页]\n{text}")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("reading DOCX knowledge files requires python-docx") from exc

    document = docx.Document(str(path))
    lines = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _normalize_text(text: str) -> str:
    lines = [" ".join(line.strip().split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)
